import os
import io
import pymupdf
import docx
import zipfile
from google import genai
from google.genai import types

def extract_text_from_pdf(file_bytes: bytes) -> str:
    text = ""
    try:
        doc = pymupdf.open(stream=file_bytes, filetype="pdf")
        for page in doc:
            text += page.get_text() + "\n"
        doc.close()
    except Exception as e:
        print(f"Error reading PDF: {e}")
    return text

def extract_text_from_docx(file_bytes: bytes) -> str:
    text = ""
    try:
        doc = docx.Document(io.BytesIO(file_bytes))
        for para in doc.paragraphs:
            text += para.text + "\n"
    except Exception as e:
        print(f"Error reading DOCX: {e}")
    return text

def extract_text_from_md(file_bytes: bytes) -> str:
    try:
        return file_bytes.decode('utf-8')
    except:
        return file_bytes.decode('latin-1')

def get_templates(template_dir="document_template"):
    templates = []
    if not os.path.exists(template_dir):
        return templates
        
    for category in sorted(os.listdir(template_dir)):
        cat_path = os.path.join(template_dir, category)
        if os.path.isdir(cat_path):
            for file in sorted(os.listdir(cat_path)):
                if file.endswith(".md"):
                    templates.append({
                        "category": category,
                        "title": file.replace(".md", ""),
                        "path": f"{category}/{file}"
                    })
    return templates

def get_template_content(template_path: str, template_dir="document_template") -> str:
    full_path = os.path.join(template_dir, template_path)
    if os.path.exists(full_path):
        with open(full_path, 'r', encoding='utf-8') as f:
            return f.read()
    return ""

def _generate_document_content(template_content: str, source_context: str, language: str = "en") -> str:
    api_key = os.environ.get("GEMINI_API_KEY")
    if not api_key:
        context_preview = source_context[:500] + ("..." if len(source_context) > 500 else "")
        return f"> **Error**: GEMINI_API_KEY environment variable not set. Please set it or add it to your environment to enable real generation.\n> Source preview: _{context_preview}_\n\n---\n\n{template_content}"
        
    try:
        lang_instruction = "English" if language == "en" else "Indonesian (Bahasa Indonesia)"
        client = genai.Client(api_key=api_key)
        prompt = f"""You are a professional technical writer and business analyst.
Your task is to generate a comprehensive document by filling in the provided template based on the source context.
Do not include any placeholders like {{{{VARIABLE}}}} in the output if they can be inferred.
If information is missing from the source context, either make a reasonable, professional assumption based on the context, or write "Information not provided in source document".

CRITICAL INSTRUCTION: You MUST generate the ENTIRE document content in {lang_instruction}. Translate any headings or content from the template into {lang_instruction} if necessary.

=== SOURCE CONTEXT ===
{source_context}

=== TEMPLATE TO FILL ===
{template_content}

Generate the final filled document in Markdown format in {lang_instruction}. Output the document content directly, do not use markdown code blocks like ```markdown...``` around the entire response.
"""
        response = client.models.generate_content(
            model='gemini-2.5-flash',
            contents=prompt,
            config=types.GenerateContentConfig(
                temperature=0.2,
            )
        )
        
        output = response.text
        if output.startswith("```markdown"):
            output = output[11:]
            if output.endswith("```"):
                output = output[:-3]
        elif output.startswith("```"):
            output = output[3:]
            if output.endswith("```"):
                output = output[:-3]
        
        return output.strip()
    except Exception as e:
        return f"> **Error generating document**: {e}\n\n---\n\n{template_content}"

def create_zip_archive(documents: list) -> io.BytesIO:
    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, "a", zipfile.ZIP_DEFLATED, False) as zip_file:
        for doc in documents:
            filename = f"{doc['category']}_{doc['title']}.md"
            zip_file.writestr(filename, doc['content'] or 'No content generated.')
    zip_buffer.seek(0)
    return zip_buffer
